from __future__ import annotations

import hashlib
import io
import json
import os
import secrets
import tempfile
import uuid
import unittest
from datetime import datetime, timezone
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader
from sqlalchemy import create_engine, text
from sqlalchemy.engine import URL

from tests.postgres_authority_fixture import seed_signed_case_authority
from tests.postgres_security_fixtures import apply_operator_security_schema
from rtm_core.authority_repository import (
    DocumentReviewAttestation,
    create_family_resolution,
    create_validated_facts,
    freeze_validated_facts,
    get_validated_facts,
    invalidate_validated_facts,
    lock_family_resolution,
    model_digest,
)
from rtm_core.contracts import (
    FactStatus,
    PreviewStatus,
    SourceReference,
    ValidatedFacts,
)
from rtm_core.family_core import resolve_family
from rtm_core.facts_review import ReviewFactsBody, review_facts
from rtm_core.generation_gateway import (
    approve_resource_for_submission,
    generate_from_frozen_preview,
)
from rtm_core.migration_router import authority_v1_ddl
from rtm_core.preview_repository import (
    approve_preview,
    create_preview,
    freeze_preview,
    submit_for_review,
)
from rtm_core.reanalysis_adapter import (
    build_validated_facts_from_reanalysis,
    load_latest_reanalysis_snapshot,
)
from rtm_core.specialist_registry import build_legal_preview


RUN_POSTGRES_INTEGRATION = os.getenv("RTM_CORE_INTEGRATION_DB") == "1"
DATABASE_URL = os.getenv("DATABASE_URL", "")


@unittest.skipUnless(
    RUN_POSTGRES_INTEGRATION and DATABASE_URL,
    "Requiere PostgreSQL temporal de RTM CORE",
)
class PostgresAuthorityIntegrationTest(unittest.TestCase):
    """Prueba el circuito real sobre PostgreSQL, sin B2 ni expedientes reales."""

    @classmethod
    def setUpClass(cls):
        cls.engine = create_engine(DATABASE_URL, pool_pre_ping=True)
        cls._reset_legacy_schema()
        cls._apply_core_migration_twice()

    @classmethod
    def tearDownClass(cls):
        cls.engine.dispose()

    def setUp(self):
        self.enterContext(
            patch.dict(
                os.environ,
                {"RTM_AUTHORITY_SIGNING_SECRET": secrets.token_urlsafe(48)},
            )
        )

    @classmethod
    def _reset_legacy_schema(cls):
        with cls.engine.begin() as conn:
            conn.execute(text("DROP SCHEMA public CASCADE"))
            conn.execute(text("CREATE SCHEMA public"))
            conn.execute(text("CREATE EXTENSION IF NOT EXISTS pgcrypto"))
            conn.execute(
                text(
                    """
                    CREATE TABLE cases (
                        id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
                        status TEXT NOT NULL DEFAULT 'uploaded',
                        payment_status TEXT,
                        authorized BOOLEAN NOT NULL DEFAULT FALSE,
                        authorized_at TIMESTAMPTZ,
                        department TEXT,
                        case_type TEXT,
                        category TEXT,
                        interested_data JSONB,
                        expediente_ref TEXT,
                        organismo TEXT,
                        contact_email TEXT,
                        test_mode BOOLEAN NOT NULL DEFAULT FALSE,
                        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                        updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                    )
                    """
                )
            )
            conn.execute(
                text(
                    """
                    CREATE TABLE documents (
                        id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
                        case_id UUID NOT NULL REFERENCES cases(id) ON DELETE CASCADE,
                        kind TEXT NOT NULL,
                        b2_bucket TEXT,
                        b2_key TEXT,
                        sha256 TEXT,
                        mime TEXT,
                        size_bytes BIGINT,
                        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                    )
                    """
                )
            )
            conn.execute(
                text(
                    """
                    CREATE TABLE extractions (
                        id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
                        case_id UUID NOT NULL REFERENCES cases(id) ON DELETE CASCADE,
                        extracted_json JSONB NOT NULL,
                        confidence DOUBLE PRECISION,
                        model TEXT,
                        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                    )
                    """
                )
            )
            conn.execute(
                text(
                    """
                    CREATE TABLE events (
                        id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
                        case_id UUID REFERENCES cases(id) ON DELETE CASCADE,
                        type TEXT NOT NULL,
                        payload JSONB,
                        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                    )
                    """
                )
            )

    @classmethod
    def _apply_core_migration_twice(cls):
        # La segunda pasada debe ser inocua: Render puede reintentar una migración.
        for _ in range(2):
            with cls.engine.begin() as conn:
                for _, statement in authority_v1_ddl():
                    conn.execute(text(statement))
                apply_operator_security_schema(conn)

        with cls.engine.begin() as conn:
            tables = {
                row[0]
                for row in conn.execute(
                    text(
                        """
                        SELECT table_name
                        FROM information_schema.tables
                        WHERE table_schema='public'
                        """
                    )
                ).fetchall()
            }
            for expected in (
                "rtm_validated_facts",
                "rtm_family_resolutions",
                "rtm_legal_previews",
                "rtm_generated_resources",
            ):
                if expected not in tables:
                    raise AssertionError(f"Falta tabla migrada: {expected}")

    def test_human_correction_versions_concurrency_and_rollback(self):
        case_id, document_id = str(uuid.uuid4()), str(uuid.uuid4())
        with self.engine.begin() as conn:
            conn.execute(text("""
                INSERT INTO cases(id,status,payment_status,authorized,department,case_type,category,interested_data,test_mode)
                VALUES (:id,'facts_validation','paid',TRUE,'traffic','fine','traffic',CAST(:identity AS JSONB),TRUE)
            """), {"id": case_id, "identity": json.dumps({"full_name": "Persona de prueba", "dni_nie": "12345678Z", "domicilio_notif": "Calle de Prueba 1, Manresa"})})
            conn.execute(text("INSERT INTO documents(id,case_id,kind) VALUES (:id,:case_id,'original')"), {"id": document_id, "case_id": case_id})
            seed_signed_case_authority(conn, case_id)
            facts = ValidatedFacts(case_id=case_id, service="traffic", extractor_version="synthetic_review_test",
                source_document_ids=[document_id], unresolved=["fecha_notificacion", "matricula"], facts={
                    "fecha_notificacion": {"status": "unresolved"}, "matricula": {"status": "unresolved"}})
            original = create_validated_facts(conn, case_id=case_id, facts=facts, created_by="operator:fixture")
        correction = ReviewFactsBody(expected_payload_sha256=original.payload_sha256, reason="Fecha contrastada con original",
            changes=[dict(field="fecha_notificacion", value="2026-09-20", document_id=document_id,
                          page_index=0, evidence="Notificación: 20/09/2026")])

        # A failure between invalidation and insert must restore the old active draft.
        with self.assertRaisesRegex(RuntimeError, "synthetic write failure"):
            with self.engine.begin() as conn, patch("rtm_core.facts_review.repository.create_validated_facts", side_effect=RuntimeError("synthetic write failure")):
                review_facts(conn, case_id=case_id, facts_id=original.id, body=correction, actor="operator:reviewer")
        with self.engine.begin() as conn:
            restored = get_validated_facts(conn, case_id, original.id)
            self.assertIsNone(restored.invalidated_at)
            self.assertEqual(restored.payload_sha256, original.payload_sha256)
            self.assertEqual(conn.execute(text("SELECT count(*) FROM rtm_validated_facts WHERE case_id=:id"), {"id": case_id}).scalar_one(), 1)
            corrected = review_facts(conn, case_id=case_id, facts_id=original.id, body=correction, actor="operator:reviewer")

        with self.engine.begin() as conn:
            stored = get_validated_facts(conn, case_id, corrected.id)
            self.assertEqual(stored.sequence, 2)
            self.assertEqual(stored.supersedes_id, original.id)
            self.assertEqual(stored.created_by, "operator:reviewer")
            self.assertFalse(stored.frozen)
            self.assertEqual(stored.facts.unresolved, ["matricula"])
            source = stored.facts.facts["fecha_notificacion"].sources[0]
            self.assertEqual((source.source_type, source.page_index, source.document_id), ("operator_document_review", 0, document_id))
            self.assertIsNotNone(get_validated_facts(conn, case_id, original.id).invalidated_at)
            event = conn.execute(text("SELECT payload FROM events WHERE case_id=:id AND type='rtm_validated_facts_reviewed'"), {"id": case_id}).scalar_one()
            self.assertEqual(event["actor"], "operator:reviewer")
            self.assertEqual(event["payload_sha256"], stored.payload_sha256)

        # A second reviewer holding the first version cannot overwrite the new one.
        with self.assertRaises(HTTPException) as stale:
            with self.engine.begin() as conn:
                review_facts(conn, case_id=case_id, facts_id=original.id, body=correction, actor="operator:other")
        self.assertEqual(stale.exception.status_code, 409)
        with self.engine.begin() as conn:
            self.assertIsNone(get_validated_facts(conn, case_id, corrected.id).invalidated_at)
            self.assertEqual(conn.execute(text("SELECT count(*) FROM rtm_validated_facts WHERE case_id=:id"), {"id": case_id}).scalar_one(), 2)

        # Even a provenance-listed document must still be an original of this case.
        with self.engine.begin() as conn:
            conn.execute(text("UPDATE documents SET kind='generated' WHERE id=:id"), {"id": document_id})
        current = correction.model_copy(update={"expected_payload_sha256": corrected.payload_sha256})
        with self.assertRaises(HTTPException) as wrong_document:
            with self.engine.begin() as conn:
                review_facts(conn, case_id=case_id, facts_id=corrected.id, body=current, actor="operator:reviewer")
        self.assertEqual(wrong_document.exception.status_code, 409)
        with self.engine.begin() as conn:
            self.assertIsNone(get_validated_facts(conn, case_id, corrected.id).invalidated_at)

    def test_added_documentary_facts_are_versioned_and_update_parking_preparation(self):
        from rtm_core.traffic_parking_preparation import build_parking_preparation
        case_id, document_id = str(uuid.uuid4()), str(uuid.uuid4())
        actor = "operator:" + str(uuid.uuid4())
        with self.engine.begin() as conn:
            conn.execute(text("""INSERT INTO cases(id,status,payment_status,authorized,department,case_type,interested_data,test_mode)
                VALUES (:id,'facts_validation','paid',TRUE,'traffic','fine','{}'::jsonb,TRUE)"""), {"id":case_id})
            conn.execute(text("INSERT INTO documents(id,case_id,kind) VALUES (:id,:case,'original')"), {"id":document_id,"case":case_id})
            seed_signed_case_authority(conn,case_id)
            facts=ValidatedFacts(case_id=case_id,service="traffic",extractor_version="synthetic_addition_test",
                source_document_ids=[document_id],facts={"hecho_denunciado_literal": {
                    "value":"Estacionar en una zona de prueba señalizada.","status":"validated","confidence":1.0,"sources":[{
                        "document_id":document_id,"page_index":0,"source_type":"operator_document_review",
                        "extraction_method":"synthetic_test","confidence":1.0,"evidence":"Estacionar en una zona de prueba señalizada."}]}})
            original=create_validated_facts(conn,case_id=case_id,facts=facts,created_by=actor)
            original_case=conn.execute(text("SELECT to_jsonb(c) FROM cases c WHERE id=:id"),{"id":case_id}).scalar_one()
        addition=ReviewFactsBody(expected_payload_sha256=original.payload_sha256,reason="Datos de prueba contrastados con el original",
            changes=[dict(operation="add",field=key,value=value,document_id=document_id,page_index=0,evidence=evidence)
                for key,value,evidence in [
                    ("lugar_infraccion","Calle ficticia 1","Lugar: Calle ficticia 1"),
                    ("pago_multa_reducido",False,"Estado: no abonada con reducción"),
                    ("fase_procedimental","notificación de denuncia","Notificación de denuncia")]])
        with self.assertRaisesRegex(RuntimeError,"addition rollback"):
            with self.engine.begin() as conn, patch("rtm_core.facts_review.repository.create_validated_facts",side_effect=RuntimeError("addition rollback")):
                review_facts(conn,case_id=case_id,facts_id=original.id,body=addition,actor=actor)
        with self.engine.begin() as conn:
            self.assertIsNone(get_validated_facts(conn,case_id,original.id).invalidated_at)
            saved=review_facts(conn,case_id=case_id,facts_id=original.id,body=addition,actor=actor)
        with self.engine.begin() as conn:
            stored=get_validated_facts(conn,case_id,saved.id)
            old=get_validated_facts(conn,case_id,original.id)
            self.assertEqual(stored.supersedes_id,original.id)
            self.assertEqual(stored.sequence,2)
            self.assertEqual(stored.created_by,actor)
            self.assertFalse(stored.frozen)
            self.assertNotIn("lugar_infraccion",old.facts.facts)
            self.assertIs(stored.facts.facts["pago_multa_reducido"].value,False)
            self.assertEqual(stored.facts.facts["lugar_infraccion"].sources[0].document_id,document_id)
            guide=build_parking_preparation(stored.facts)
            self.assertIn("no se ha pagado",guide["payment_note"])
            self.assertIn("Consta una fase inicial",guide["procedure_note"])
            self.assertFalse(guide["ready_to_submit"])
            event=conn.execute(text("SELECT payload FROM events WHERE case_id=:id AND type='rtm_validated_facts_reviewed'"),{"id":case_id}).scalar_one()
            self.assertEqual(event["added_fields"],["fase_procedimental","lugar_infraccion","pago_multa_reducido"])
            self.assertEqual(event["corrected_fields"],[])
            self.assertEqual(event["actor"],actor)
            self.assertEqual(event["payload_sha256"],stored.payload_sha256)
            current_case=conn.execute(text("SELECT to_jsonb(c) FROM cases c WHERE id=:id"),{"id":case_id}).scalar_one()
            self.assertEqual({key:value for key,value in current_case.items() if key != "updated_at"},
                             {key:value for key,value in original_case.items() if key != "updated_at"})
        # Neither a stale retry nor an explicit add against the current version
        # may overwrite the already recorded data.
        for record_id, expected in [(original.id,original.payload_sha256),(saved.id,saved.payload_sha256)]:
            with self.assertRaises(HTTPException):
                with self.engine.begin() as conn:
                    review_facts(conn,case_id=case_id,facts_id=record_id,
                        body=addition.model_copy(update={"expected_payload_sha256":expected}),actor=actor)
        with self.engine.begin() as conn:
            self.assertIsNone(get_validated_facts(conn,case_id,saved.id).invalidated_at)
            self.assertEqual(conn.execute(text("SELECT count(*) FROM rtm_validated_facts WHERE case_id=:id"),{"id":case_id}).scalar_one(),2)
            self.assertEqual(conn.execute(text("SELECT count(*) FROM rtm_generated_resources WHERE case_id=:id"),{"id":case_id}).scalar_one(),0)

    def _ops_case_detail(self, case_id):
        import ops_operator_router as router
        from fastapi import Request
        from types import SimpleNamespace
        scope = SimpleNamespace(individual_session=True, scope_all=True, operator_id=str(uuid.uuid4()))
        with patch.object(router, "get_engine", return_value=self.engine), \
             patch.object(router, "require_operator_token"), \
             patch.object(router, "load_ops_case_scope", return_value=scope), \
             patch.object(router, "require_case_in_scope", return_value=case_id):
            # Only authentication is replaced; SQL and cryptographic verification are real.
            return router.get_case_detail(case_id, Request({"type": "http"}), "fixture-token")

    def test_local_fine_fixture_is_new_idempotent_and_never_approves_signature(self):
        from scripts import rtm_local_fine_fixture as fixture
        from scripts import rtm_local_candidate_repair as repair
        from case_authority import verify_signed_case_authority, verify_authorization_signature_candidate
        import ops_operator_router as router
        from fastapi import Request
        from types import SimpleNamespace
        document_root = self.enterContext(tempfile.TemporaryDirectory(prefix="rtm-candidate-check-"))
        # Storage's local profile is synthetic; every SQL call uses self.engine,
        # already bound to the runner's verified ephemeral PostgreSQL cluster.
        self.enterContext(patch.dict(os.environ, {
            "RTM_ENV": "development", "RTM_LOCAL_BIND_HOST": "127.0.0.1",
            "RTM_ALLOWED_HOSTS": "127.0.0.1", "ALLOWED_ORIGINS": "http://127.0.0.1:5173",
            "RTM_ENABLE_LOCAL_OPERATOR_AUTH": "1", "RTM_ENABLE_OPERATOR_AUTH_V1": "1",
            "RTM_ENABLE_LOCAL_DOCUMENT_STORAGE": "1", "RTM_LOCAL_DOCUMENT_ROOT": document_root,
            "RTM_TRUST_PROXY_HEADERS": "0", "RTM_ALLOW_REAL_CUSTOMER_DATA": "0",
            "RTM_ENABLE_B2": "0", "RTM_ENABLE_STRIPE": "0",
            "RTM_ENABLE_FINAL_PAYMENTS": "0", "RTM_ENABLE_DOCUMENT_PROVIDER": "0",
            "RTM_ENABLE_OUTBOUND_EMAIL": "0", "RTM_ENABLE_EXTERNAL_SUBMISSION": "0",
            "DATABASE_URL": URL.create("postgresql+psycopg", username="rtm_local_app",
                password="unused_test_value", host="127.0.0.1", port=5432,
                database="rtm_local").render_as_string(hide_password=False),
        }))
        objects = {}
        def store(case_id, kind, data, extension, mime):
            self.assertEqual(case_id, fixture.CASE_ID)
            self.assertEqual(mime, "application/pdf")
            content = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
            self.assertIn("FICTICI", content)
            self.assertIn("RTMTEST002", content)
            objects[kind] = data
            return fixture.storage.upload_bytes(case_id, kind, data, extension, mime)
        def require_ephemeral_database(conn):
            # Local runs verify their own cluster; GitHub uses the disposable
            # PostgreSQL service configured in rtm-core-ci.yml.
            from sqlalchemy.engine import make_url
            expected = make_url(DATABASE_URL)
            actual = conn.execute(text("SELECT current_database(),current_user,inet_server_port()" )).one()
            github_test_service = (
                os.getenv("GITHUB_ACTIONS") == "true"
                and (expected.host, expected.port, expected.username, expected.database)
                == ("127.0.0.1", 5432, "postgres", "rtm_core_test")
            )
            self.assertTrue(expected.database.startswith("rtm_check_") or github_test_service)
            self.assertEqual(tuple(actual), (expected.database, expected.username, expected.port))
        with self.engine.begin() as conn:
            for name, kind in [("contact_name", "TEXT"), ("customer_comment", "TEXT"),
                               ("source_module", "TEXT"), ("product_code", "TEXT"), ("paid_at", "TIMESTAMPTZ")]:
                conn.execute(text(f"ALTER TABLE cases ADD COLUMN IF NOT EXISTS {name} {kind}"))
            from ops_followups_schema import ops_followups_ddl
            for _, statement in ops_followups_ddl():
                conn.execute(text(statement))
        def candidate_response():
            operator_id, session_id = str(uuid.uuid4()), str(uuid.uuid4())
            scope = SimpleNamespace(individual_session=True, role_code="rtm.supervisor",
                                    permissions=("ops.view", "ops.supervise"))
            request = Request({"type": "http"})
            request.state.rtm_operator_context = SimpleNamespace(operator_id=operator_id,
                session_id=session_id, actor="operator:" + operator_id)
            with patch.object(router, "get_engine", return_value=self.engine), \
                 patch.object(router, "require_operator_token"), \
                 patch.object(router, "load_ops_case_scope", return_value=scope), \
                 patch.object(router, "require_case_in_scope", return_value=fixture.CASE_ID):
                # Only authentication is replaced. Real route, SQL, storage, PDF
                # validation, hash checks, signed view event and response bytes.
                return router.view_authorization_signature_candidate(fixture.CASE_ID,
                    fixture.document_id("authorization_signed_candidate"), request, "fixture-token")

        with patch.object(fixture, "require_local_database", side_effect=require_ephemeral_database), \
             patch.object(repair, "require_local_database", side_effect=require_ephemeral_database):
            with self.engine.begin() as conn:
                before = conn.execute(text("SELECT count(*) FROM cases")).scalar_one()
                result = fixture.seed_case(conn, store)
                self.assertTrue(result["created"])
            detail = self._ops_case_detail(fixture.CASE_ID)
            self.assertEqual(detail["case_type"], "fine")
            self.assertEqual(detail["authorization_evidence_status"], "pending_review")
            self.assertFalse(detail["signed_authority_verified"])
            response = candidate_response()
            original_pdf = objects["authorization_signature_candidate"]
            self.assertEqual(response.body, original_pdf)
            self.assertIn("no-store", response.headers["cache-control"])
            candidate_id = fixture.document_id("authorization_signed_candidate")
            # Reproduce the original fixture's wrong folder and exact HTTP 409.
            old_bucket, old_key = fixture.storage.upload_bytes(fixture.CASE_ID,
                "authorization_signed_candidate", original_pdf, ".pdf", "application/pdf")
            with self.engine.begin() as conn:
                conn.execute(text("UPDATE documents SET b2_key=:key WHERE id=:id"),
                             {"key": old_key, "id": candidate_id})
            with self.assertRaises(HTTPException) as wrong_folder:
                candidate_response()
            self.assertEqual(wrong_folder.exception.status_code, 409)
            self.assertEqual(wrong_folder.exception.detail, "Candidato de firma no verificable")

            # A modified document must be refused, even in this synthetic repair.
            _, tampered_key = fixture.storage.upload_bytes(fixture.CASE_ID,
                "authorization_signed_candidate", original_pdf + b"changed", ".pdf", "application/pdf")
            with self.engine.begin() as conn:
                conn.execute(text("UPDATE documents SET b2_key=:key WHERE id=:id"),
                             {"key": tampered_key, "id": candidate_id})
            with self.assertRaisesRegex(RuntimeError, "huella"):
                with self.engine.begin() as conn:
                    repair.repair_candidate(conn)
            with self.engine.begin() as conn:
                conn.execute(text("UPDATE documents SET b2_key=:key WHERE id=:id"),
                             {"key": old_key, "id": candidate_id})

            # Failure after the UPDATE must roll back both relocation and event.
            with self.assertRaisesRegex(RuntimeError, "synthetic audit failure"):
                with self.engine.begin() as conn, patch.object(fixture, "event",
                        side_effect=RuntimeError("synthetic audit failure")):
                    repair.repair_candidate(conn)
            with self.engine.begin() as conn:
                self.assertEqual(conn.execute(text("SELECT b2_key FROM documents WHERE id=:id"),
                    {"id": candidate_id}).scalar_one(), old_key)
                chain_before = verify_authorization_signature_candidate(conn, fixture.CASE_ID, candidate_id)
                repaired = repair.repair_candidate(conn)
                self.assertTrue(repaired["changed"])
                self.assertTrue(repaired["pdf_download_verified"])
                self.assertEqual(chain_before,
                    verify_authorization_signature_candidate(conn, fixture.CASE_ID, candidate_id))
            self.assertEqual(candidate_response().body, original_pdf)
            self.assertEqual(fixture.storage.download_bytes_limited(old_bucket, old_key,
                max_bytes=repair.MAX_BYTES, case_id=fixture.CASE_ID), original_pdf)
            with self.engine.begin() as conn, patch.object(fixture.storage, "upload_bytes") as upload:
                repeated = repair.repair_candidate(conn)
                self.assertFalse(repeated["changed"])
                upload.assert_not_called()
                self.assertEqual(conn.execute(text("SELECT count(*) FROM events WHERE case_id=:id AND type=:kind"),
                    {"id": fixture.CASE_ID, "kind": repair.REPAIR_EVENT}).scalar_one(), 1)
            with self.engine.begin() as conn:
                self.assertEqual(conn.execute(text("SELECT count(*) FROM cases")).scalar_one(), before + 1)
                row = conn.execute(text("SELECT payment_status,product_code,test_mode,interested_data FROM cases WHERE id=:id"), {"id": fixture.CASE_ID}).one()
                self.assertEqual(row[0:3], ("paid", "RTM_LOCAL_SIMULATED_REVIEW", True))
                self.assertTrue(row[3]["local_test"]["payment_simulated"])
                self.assertFalse(row[3]["local_test"]["real_charge"])
                with self.assertRaises(HTTPException) as pending:
                    verify_signed_case_authority(conn, fixture.CASE_ID)
                self.assertEqual(pending.exception.status_code, 409)
                self.assertEqual(conn.execute(text("SELECT count(*) FROM events WHERE case_id=:id AND type='authorization_signature_approved'"), {"id": fixture.CASE_ID}).scalar_one(), 0)
                stored = get_validated_facts(conn, fixture.CASE_ID, result["facts_id"])
                self.assertFalse(stored.frozen)
                self.assertEqual(len(stored.facts.unresolved), 7)
                self.assertTrue(all(fact.value is None for fact in stored.facts.facts.values()))
            with self.engine.begin() as conn:
                existing = fixture.seed_case(conn, lambda *args: self.fail("Duplicate fixture must not write documents"))
                self.assertFalse(existing["created"])
                self.assertEqual(conn.execute(text("SELECT count(*) FROM documents WHERE case_id=:id"), {"id": fixture.CASE_ID}).scalar_one(), 5)
            # Explicitly simulated filing: retained receipt and initiation notice,
            # persistent clocks, no real submission/authorization/payment changes.
            from scripts import rtm_local_submission_fixture as filing
            from rtm_core.post_filing_deadlines import local_post_filing_projection
            with patch.object(filing, "require_local_database", side_effect=require_ephemeral_database):
                with self.engine.begin() as conn:
                    before_case = conn.execute(text("SELECT to_jsonb(c) FROM cases c WHERE id=:id"),
                                               {"id": fixture.CASE_ID}).scalar_one()
                    simulated = filing.seed_submission(conn)
                    self.assertTrue(simulated["created"])
                    self.assertEqual(simulated["status"], "presented_simulated")
                    self.assertEqual(simulated["calculation"]["reference_due_on"], "2027-09-17")
                    self.assertIsNone(simulated["calculation"]["legal_due_on"])
                    self.assertEqual(simulated["followups"]["created_count"], 2)
                # Reproduce the user's existing filing with no inbox reminders.
                with self.engine.begin() as conn:
                    conn.execute(text("DELETE FROM ops_followups WHERE case_id=:id"), {"id": fixture.CASE_ID})
                # The audit and both reminder inserts must commit atomically.
                with self.assertRaisesRegex(RuntimeError, "synthetic audit failure"):
                    with self.engine.begin() as conn, patch.object(fixture, "event",
                            side_effect=RuntimeError("synthetic audit failure")):
                        filing.seed_submission(conn)
                with self.engine.begin() as conn:
                    self.assertEqual(conn.execute(text("SELECT count(*) FROM ops_followups WHERE case_id=:id"),
                        {"id": fixture.CASE_ID}).scalar_one(), 0)
                    backfilled = filing.seed_submission(conn)
                    self.assertFalse(backfilled["created"])
                    self.assertEqual(backfilled["followups"]["created_count"], 2)
                    self.assertEqual(backfilled["followups"]["items"], simulated["followups"]["items"])
                # Exercise the existing inbox route and its real SQL. Authentication
                # alone uses a supervisor scope; scope security has its own suite.
                import ops
                from datetime import date, timedelta
                from rtm_core.post_filing_deadlines import MADRID
                scope = SimpleNamespace(individual_session=False, scope_all=True)
                with patch.object(ops, "get_engine", return_value=self.engine), \
                     patch.object(ops, "_require_operator"), \
                     patch.object(ops, "load_ops_case_scope", return_value=scope):
                    inbox = ops.list_all_followups(Request({"type": "http"}), "fixture-token", "pending", 500)
                rows = [item for item in inbox["items"] if item["case_id"] == fixture.CASE_ID]
                self.assertEqual(len(rows), 2)
                self.assertTrue(all(item["title"].startswith("PRUEBA LOCAL") for item in rows))
                self.assertTrue(all(item["expediente_ref"] == fixture.REFERENCE for item in rows))
                self.assertEqual(rows[0]["due_at"].astimezone(MADRID).date(),
                    date.fromisoformat(simulated["calculation"]["filing_date"]) + timedelta(days=1))
                self.assertEqual(rows[1]["due_at"].astimezone(MADRID).date(), date(2027, 9, 10))
                with self.engine.begin() as conn:
                    conn.execute(text("UPDATE ops_followups SET status='resolved',resolved_at=NOW(),"
                        "resolved_by='operator:test',resolution_note='Reviewed' WHERE id=:id"), {"id": rows[0]["id"]})
                    followups_before = conn.execute(text("SELECT to_jsonb(f) FROM ops_followups f WHERE case_id=:id ORDER BY id"),
                                                   {"id": fixture.CASE_ID}).scalars().all()
                with self.engine.begin() as conn:
                    self.assertEqual(conn.execute(text("SELECT to_jsonb(c) FROM cases c WHERE id=:id"),
                                                  {"id": fixture.CASE_ID}).scalar_one(), before_case)
                    repeated = filing.seed_submission(conn)
                    self.assertFalse(repeated["created"])
                    self.assertEqual(repeated["followups"]["created_count"], 0)
                    self.assertEqual(conn.execute(text("SELECT to_jsonb(f) FROM ops_followups f WHERE case_id=:id ORDER BY id"),
                        {"id": fixture.CASE_ID}).scalars().all(), followups_before)
                    self.assertEqual(repeated["event_id"], simulated["event_id"])
                    self.assertEqual(conn.execute(text("SELECT count(*) FROM documents WHERE case_id=:id"),
                                                  {"id": fixture.CASE_ID}).scalar_one(), 7)
                    self.assertEqual(conn.execute(text("SELECT count(*) FROM events WHERE case_id=:id AND type IN ('authorization_signature_approved','manual_submission_registered','dgt_submitted')"),
                                                  {"id": fixture.CASE_ID}).scalar_one(), 0)
                    conn.execute(text("UPDATE documents SET sha256=:hash WHERE id=:id"),
                                 {"id": simulated["receipt_document_id"], "hash": "0" * 64})
                    blocked = local_post_filing_projection(conn, fixture.CASE_ID)
                    self.assertEqual(blocked["status"], "evidence_unverifiable")
                    self.assertNotIn("calculation", blocked)
                    with self.assertRaisesRegex(RuntimeError, "necesita revision"):
                        filing.seed_submission(conn)
                    conn.execute(text("UPDATE documents SET sha256=:hash WHERE id=:id"),
                                 {"id": simulated["receipt_document_id"], "hash": simulated["receipt_sha256"]})
                detail = self._ops_case_detail(fixture.CASE_ID)
                self.assertEqual(detail["post_filing"]["status"], "presented_simulated")
                self.assertEqual(detail["post_filing"]["calculation"]["reference_due_on"], "2027-09-17")
                # Human deadline review persists separately from the filing and
                # never grants signature authority or a procedural outcome.
                from rtm_core import post_filing_review as deadline_review
                from tests.test_post_filing_review import body as deadline_body
                operator_id = str(uuid.uuid4())
                review_request = Request({"type": "http"})
                review_request.state.rtm_operator_context = SimpleNamespace(operator_id=operator_id,
                    session_id=str(uuid.uuid4()), actor="operator:" + operator_id)
                review_scope = SimpleNamespace(individual_session=True, scope_all=True, operator_id=operator_id,
                    role_code="rtm.supervisor", permissions=("ops.view", "ops.supervise"))
                def save_clock(projection, **changes):
                    review = projection["review"]
                    body = deadline_review.DeadlineReviewBody(**deadline_body(
                        expected_source_sha256=review["source_sha256"], expected_review_id=review["latest_id"], **changes))
                    with patch.object(router, "get_engine", return_value=self.engine), \
                         patch.object(router, "require_operator_token"), \
                         patch.object(router, "load_ops_case_scope", return_value=review_scope):
                        return router.review_case_deadlines(fixture.CASE_ID, body, review_request, "fixture-token")["post_filing"]
                initial = detail["post_filing"]
                pending_review = save_clock(initial)
                self.assertIsNone(pending_review["calculation"]["legal_due_on"])
                self.assertEqual(len(pending_review["review"]["history"]), 1)
                self.assertEqual(pending_review["review"]["history"][0]["actor"], "operator:" + operator_id)
                with self.assertRaises(HTTPException) as stale:
                    save_clock(initial)
                self.assertEqual(stale.exception.status_code, 409)
                # If the final projection cannot verify a new event, rollback it.
                original_projection = local_post_filing_projection
                calls = []
                def fail_after_insert(conn, case_id):
                    calls.append(case_id)
                    return original_projection(conn, case_id) if len(calls) == 1 else {"status": "evidence_unverifiable"}
                with patch("rtm_core.post_filing_deadlines.local_post_filing_projection", side_effect=fail_after_insert), \
                     self.assertRaisesRegex(RuntimeError, "verificar"):
                    save_clock(pending_review)
                self.assertEqual(self._ops_case_detail(fixture.CASE_ID)["post_filing"]["review"]["latest_id"],
                                 pending_review["review"]["latest_id"])
                calendar = dict(from_date="2027-09-01", to_date="2027-09-30", holidays=["2027-09-17"],
                    source="Calendario inventado exclusivamente para prueba de desplazamiento", territory="Municipio ficticio RTM")
                complete = save_clock(pending_review, calendar=calendar, rule_checked=True, procedural_status="no_changes")
                self.assertEqual(complete["calculation"]["legal_due_on"], "2027-09-20")
                self.assertFalse(complete["calculation"]["automatic_legal_consequence"])
                reloaded = self._ops_case_detail(fixture.CASE_ID)["post_filing"]
                self.assertEqual(reloaded["calculation"]["legal_due_on"], "2027-09-20")
                self.assertEqual(len(reloaded["review"]["history"]), 2)
                # New case evidence makes the old review stale instead of keeping
                # a previously confirmed date active indefinitely.
                with self.engine.begin() as conn:
                    fixture.event(conn, "local_test_procedural_change", {"synthetic": True}, datetime.now(timezone.utc))
                changed = self._ops_case_detail(fixture.CASE_ID)["post_filing"]
                self.assertTrue(changed["review"]["stale"])
                self.assertIsNone(changed["calculation"]["legal_due_on"])
                with self.assertRaises(HTTPException) as stale_evidence:
                    save_clock(complete, calendar=calendar, rule_checked=True, procedural_status="no_changes")
                self.assertEqual(stale_evidence.exception.status_code, 409)
                needs_assessment = save_clock(changed, calendar=calendar, rule_checked=True, procedural_status="has_changes")
                self.assertIsNone(needs_assessment["calculation"]["legal_due_on"])
                self.assertIn("incidencias", " ".join(needs_assessment["calculation"]["review_reasons"]))
                with self.engine.begin() as conn:
                    self.assertEqual(conn.execute(text("SELECT to_jsonb(c) FROM cases c WHERE id=:id"),
                        {"id": fixture.CASE_ID}).scalar_one(), before_case)
                    latest_review = needs_assessment["review"]["latest_id"]
                    envelope = conn.execute(text("SELECT payload FROM events WHERE id=:id"), {"id": latest_review}).scalar_one()
                    conn.execute(text("UPDATE events SET payload='{}'::jsonb WHERE id=:id"), {"id": latest_review})
                self.assertEqual(self._ops_case_detail(fixture.CASE_ID)["post_filing"]["status"], "evidence_unverifiable")
                with self.assertRaises(HTTPException):
                    save_clock(needs_assessment)
                with self.engine.begin() as conn:
                    conn.execute(text("UPDATE events SET payload=CAST(:value AS JSONB) WHERE id=:id"),
                                 {"id": latest_review, "value": json.dumps(envelope)})
            # A colliding real/unmarked case must never become a paid fixture.
            with self.engine.begin() as conn:
                conn.execute(text("UPDATE cases SET test_mode=FALSE WHERE id=:id"), {"id": fixture.CASE_ID})
            with self.assertRaises(RuntimeError):
                with self.engine.begin() as conn:
                    fixture.seed_case(conn, lambda *args: self.fail("Collision must not write"))
            with self.assertRaisesRegex(RuntimeError, "prueba local prevista"):
                with self.engine.begin() as conn:
                    repair.repair_candidate(conn)
            with patch.object(filing, "require_local_database", side_effect=require_ephemeral_database):
                with self.assertRaisesRegex(RuntimeError, "ficticio esperado"):
                    with self.engine.begin() as conn:
                        filing.seed_submission(conn)
        self.assertEqual(len(objects), 5)

    def test_full_authority_chain_reaches_approved_resource(self):
        case_id = str(uuid.uuid4())
        document_id = str(uuid.uuid4())
        reanalysis_run_id = str(uuid.uuid4())
        wrapper = {
            "reanalysis_run_id": reanalysis_run_id,
            "storage": {"source_document_ids": [document_id]},
            "pages": [
                {
                    "page_index": 1,
                    "document_id": document_id,
                    "mime_detected": "image/tiff",
                }
            ],
            "extracted": {
                "extractor_version": "traffic_fine_reanalysis_v1_18",
                "source_document_ids": [document_id],
                # Error legacy deliberado: el adaptador no puede promoverlo.
                "familia_resuelta": "velocidad",
                "tipo_infraccion": "velocidad",
                "raw_text_blob": "CASILLA IMPRESA km/h",
                "velocidad_medida_kmh": 63,
                "velocidad_limite_kmh": 11,
                "hecho_denunciado_literal": (
                    "Conducir de forma temeraria creando un riesgo grave."
                ),
            },
        }
        event_payload = {
            "reanalysis_run_id": reanalysis_run_id,
            "extractor_version": "traffic_fine_reanalysis_v1_18",
            "handwritten_precision_version": "traffic_handwritten_precision_v1_0",
            "handwritten_precision_values": {
                "hecho_denunciado_literal": (
                    "Conducir de forma temeraria creando un riesgo grave."
                )
            },
            "handwritten_precision_confidence": {
                "hecho_denunciado_literal": 0.99,
            },
            "handwritten_precision_evidence": {
                "hecho_denunciado_literal": "CONDUCIR DE FORMA TEMERARIA",
            },
            "handwritten_precision_quality": {
                "legibility": "high",
                "legibility_score": 0.99,
            },
            "traffic_generic_facts_version": "traffic_generic_facts_v1_2",
            "traffic_generic_facts": {
                "organismo": "Servei Català de Trànsit",
                "expediente_ref": "02510067072-0",
                "document_type": "denuncia",
                "procedural_stage_hint": "initial_notice",
                "sancion_ordinaria_eur": 500,
                "puntos_detraccion": 6,
                "fecha_limite": "2026-08-20",
            },
            "traffic_generic_facts_confidence": {
                "organismo": 0.99,
                "expediente_ref": 0.99,
                "document_type": 0.98,
                "procedural_stage_hint": 0.98,
                "sancion_ordinaria_eur": 0.98,
                "puntos_detraccion": 0.98,
                "fecha_limite": 0.99,
            },
            "traffic_generic_facts_evidence": {
                "organismo": "SERVEI CATALÀ DE TRÀNSIT",
                "expediente_ref": "02510067072-0",
                "document_type": "DENÚNCIA / INICIACIÓ",
                "procedural_stage_hint": "NOTIFICACIÓ DE DENÚNCIA",
                "sancion_ordinaria_eur": "500,00 EUR",
                "puntos_detraccion": "6 PUNTS",
                "fecha_limite": "20-08-2026",
            },
            "unresolved_critical_fields": [],
            "missing_required_fields": ["fecha_notificacion"],
            "critical_conflicts_resolved": [],
        }

        with self.engine.begin() as conn:
            conn.execute(
                text(
                    """
                    INSERT INTO cases(
                        id, status, payment_status, authorized, department,
                        case_type, category, interested_data, expediente_ref,
                        organismo, contact_email, test_mode, created_at, updated_at
                    ) VALUES (
                        CAST(:case_id AS UUID), 'core_review_pending', 'paid', TRUE,
                        'traffic', 'fine', 'traffic', CAST(:interested AS JSONB),
                        '02510067072-0', 'Servei Català de Trànsit',
                        'test@example.invalid', FALSE, NOW(), NOW()
                    )
                    """
                ),
                {
                    "case_id": case_id,
                    "interested": (
                        '{"full_name":"Persona de prueba",'
                        '"dni_nie":"12345678Z",'
                        '"domicilio_notif":"Calle de Prueba 1, Manresa"}'
                    ),
                },
            )
            conn.execute(
                text(
                    """
                    INSERT INTO documents(
                        id, case_id, kind, b2_bucket, b2_key, mime,
                        size_bytes, created_at
                    ) VALUES (
                        CAST(:document_id AS UUID), CAST(:case_id AS UUID),
                        'original', 'ci', 'original/manuscrito.tif',
                        'image/tiff', 1024, NOW()
                    )
                    """
                ),
                {"document_id": document_id, "case_id": case_id},
            )
            conn.execute(
                text(
                    """
                    INSERT INTO extractions(
                        case_id, extracted_json, confidence, model, created_at
                    ) VALUES (
                        CAST(:case_id AS UUID), CAST(:payload AS JSONB), 0.99,
                        'rtm_intelligence_core_v1+traffic_fine+v1_18', NOW()
                    )
                    """
                ),
                {"case_id": case_id, "payload": json.dumps(wrapper)},
            )
            conn.execute(
                text(
                    """
                    INSERT INTO events(case_id, type, payload, created_at)
                    VALUES (
                        CAST(:case_id AS UUID), 'case_reanalysis_completed',
                        CAST(:payload AS JSONB), NOW()
                    )
                    """
                ),
                {"case_id": case_id, "payload": json.dumps(event_payload)},
            )
            seed_signed_case_authority(conn, case_id)

        detail = self._ops_case_detail(case_id)
        self.assertEqual(detail["authorization_evidence_status"], "verified")
        self.assertTrue(detail["signed_authority_verified"])

        with self.engine.begin() as conn:
            stored_wrapper, stored_event = load_latest_reanalysis_snapshot(conn, case_id)
            adapted = build_validated_facts_from_reanalysis(
                case_id=case_id,
                wrapper=stored_wrapper,
                event_payload=stored_event,
            )
            self.assertEqual(
                adapted.facts.facts["hecho_denunciado_literal"].status,
                FactStatus.UNRESOLVED,
            )
            self.assertEqual(
                adapted.facts.facts["velocidad_medida_kmh"].status,
                FactStatus.UNRESOLVED,
            )

            # Un dato ausente debe persistir como pendiente, nunca desaparecer
            # ni permitir congelar la lectura de IA sin revisión documental.
            unreviewed = create_validated_facts(
                conn, case_id=case_id, facts=adapted.facts,
                created_by="ci:unreviewed-reanalysis",
            )
            with self.assertRaises(HTTPException) as blocked:
                freeze_validated_facts(conn, case_id, unreviewed.id, "ops:ci")
            self.assertEqual(blocked.exception.status_code, 409)
            self.assertEqual(blocked.exception.detail["code"], "document_review_attestation_required")
            persisted = get_validated_facts(conn, case_id, unreviewed.id)
            self.assertFalse(persisted.frozen)
            self.assertIn("fecha_notificacion", persisted.facts.unresolved)
            self.assertIsNone(persisted.facts.facts["fecha_notificacion"].value)
            invalidate_validated_facts(
                conn, case_id, unreviewed.id, "ops:ci",
                "Sustituir la lectura candidata por la versión contrastada en la prueba.",
            )

            reviewed_payload = adapted.facts.model_dump(mode="python")
            reviewed_values = {
                "hecho_denunciado_literal": (
                    "Conducir de forma temeraria creando un riesgo grave."
                ),
                "organismo": "Servei Català de Trànsit",
                "expediente_ref": "02510067072-0",
                "tipo_documento": "denuncia",
                "fase_procedimental": "initial_notice",
                "sancion_importe_eur": 500,
                "puntos_detraccion": 6,
                "fecha_limite": "2026-08-20",
            }
            for fact_key, fact_value in reviewed_values.items():
                reviewed_payload["facts"][fact_key] = {
                    "value": fact_value,
                    "status": FactStatus.VALIDATED,
                    "confidence": 1.0,
                    "sources": [
                        source.model_dump(mode="python")
                        for source in adapted.facts.facts[fact_key].sources
                    ]
                    + [
                        SourceReference(
                            document_id=document_id,
                            page_index=0,
                            source_type="operator_document_review",
                            extraction_method="ops_document_review_v1",
                            evidence=str(fact_value),
                            confidence=1.0,
                        ).model_dump(mode="python")
                    ],
                    "conflicts": [],
                    "notes": ["Hecho contrastado con el documento por OPS."],
                }
            reviewed_payload["unresolved"] = [
                key
                for key in reviewed_payload["unresolved"]
                if key not in reviewed_values
            ]
            reviewed_facts = ValidatedFacts.model_validate(reviewed_payload)

            facts_record = create_validated_facts(
                conn,
                case_id=case_id,
                facts=reviewed_facts,
                created_by="ci:reanalysis-adapter",
                supersedes_id=unreviewed.id,
            )
            facts_record = freeze_validated_facts(
                conn,
                case_id,
                facts_record.id,
                "ops:ci",
                document_review_attestation=DocumentReviewAttestation(
                    documents_reviewed=True,
                    facts_reviewed=True,
                    source_document_ids=list(reviewed_facts.source_document_ids),
                    facts_payload_sha256=model_digest(reviewed_facts),
                    review_notes="Revisión documental completa en integración CI.",
                ),
            )

            resolution = resolve_family(facts_record.facts)
            self.assertEqual(resolution.family, "temeraria")
            self.assertNotEqual(resolution.family, "velocidad")

            family_record = create_family_resolution(
                conn,
                case_id=case_id,
                resolution=resolution,
                created_by="ci:family-core",
                validated_facts_id=facts_record.id,
            )
            family_record = lock_family_resolution(
                conn,
                case_id,
                family_record.id,
                "ops:ci",
            )

            draft = build_legal_preview(facts_record, family_record)
            self.assertFalse(
                [
                    item.code
                    for item in draft.missing_items
                    if item.severity.value == "blocking"
                ]
            )
            preview_record = create_preview(
                conn,
                case_id=case_id,
                preview=draft,
                created_by="ci:traffic.temeraria",
            )
            preview_record = submit_for_review(
                conn,
                case_id,
                preview_record.id,
                "ops:ci",
            )
            preview_record = approve_preview(
                conn,
                case_id,
                preview_record.id,
                "ops:ci",
            )
            preview_record = freeze_preview(
                conn,
                case_id,
                preview_record.id,
                "ops:ci",
            )
            self.assertEqual(preview_record.status, PreviewStatus.FROZEN)

            generated = {}

            def capture_document(upload_case_id, folder, data, extension, mime):
                self.assertEqual(upload_case_id, case_id)
                self.assertEqual(folder, "rtm_generated")
                self.assertNotIn(extension, generated)
                generated[extension] = data
                return "ci", f"generated/recurso{extension}"

            # Los dos documentos se generan de verdad; solo su custodia externa
            # se sustituye por memoria para mantener la prueba aislada.
            with patch("rtm_core.generation_gateway.upload_bytes", side_effect=capture_document):
                resource = generate_from_frozen_preview(
                    conn,
                    case_id=case_id,
                    preview_id=preview_record.id,
                    generated_by="ci:generate",
                )

            pdf = PdfReader(io.BytesIO(generated[".pdf"]))
            self.assertGreater(len(pdf.pages), 0)
            pdf_text = " ".join(" ".join(page.extract_text() for page in pdf.pages).split())
            word = Document(io.BytesIO(generated[".docx"]))
            word_text = " ".join(" ".join(item.text for item in word.paragraphs).split())
            for rendered in (pdf_text, word_text):
                self.assertIn("02510067072-0", rendered)
                self.assertIn("Persona de prueba", rendered)
                self.assertIn("conducir de forma temeraria", rendered.lower())
                self.assertNotIn("63 km/h", rendered)
            for document_id, extension in ((resource.pdf_document_id, ".pdf"),
                                           (resource.docx_document_id, ".docx")):
                stored_hash, stored_size = conn.execute(
                    text("SELECT sha256, size_bytes FROM documents WHERE id=CAST(:id AS UUID)"),
                    {"id": document_id},
                ).one()
                self.assertEqual(stored_hash, hashlib.sha256(generated[extension]).hexdigest())
                self.assertEqual(stored_size, len(generated[extension]))
            prepared_status = conn.execute(
                text("SELECT status FROM cases WHERE id=CAST(:id AS UUID)"), {"id": case_id},
            ).scalar_one()
            self.assertEqual(prepared_status, "final_ready")
            self.assertIsNone(resource.approved_at)

            resource = approve_resource_for_submission(
                conn,
                case_id=case_id,
                resource_id=resource.id,
                approved_by="ops:ci",
            )
            self.assertEqual(resource.status, "final_ready")
            self.assertEqual(resource.approved_by, "ops:ci")

            status = conn.execute(
                text("SELECT status FROM cases WHERE id=CAST(:id AS UUID)"),
                {"id": case_id},
            ).scalar_one()
            self.assertEqual(status, "ready_to_submit")

            links = conn.execute(
                text(
                    """
                    SELECT fr.validated_facts_id, lp.validated_facts_id,
                           lp.family_resolution_id, gr.legal_preview_id,
                           gr.pdf_document_id, gr.docx_document_id
                    FROM rtm_family_resolutions fr
                    JOIN rtm_legal_previews lp
                      ON lp.family_resolution_id=fr.id
                    JOIN rtm_generated_resources gr
                      ON gr.legal_preview_id=lp.id
                    WHERE fr.case_id=CAST(:case_id AS UUID)
                    """
                ),
                {"case_id": case_id},
            ).one()
            self.assertEqual(str(links[0]), facts_record.id)
            self.assertEqual(str(links[1]), facts_record.id)
            self.assertEqual(str(links[2]), family_record.id)
            self.assertEqual(str(links[3]), preview_record.id)
            self.assertTrue(links[4])
            self.assertTrue(links[5])

            # Generate es idempotente para una misma previa congelada.
            with patch("rtm_core.generation_gateway.upload_bytes") as second_upload:
                same = generate_from_frozen_preview(
                    conn,
                    case_id=case_id,
                    preview_id=preview_record.id,
                    generated_by="ci:generate-retry",
                )
                second_upload.assert_not_called()
            self.assertEqual(same.id, resource.id)

        # La revisión versionada y los pendientes sobreviven al commit y a una
        # nueva conexión, no solo a los objetos Python devueltos en memoria.
        with self.engine.connect() as conn:
            reloaded = get_validated_facts(conn, case_id, facts_record.id)
            self.assertTrue(reloaded.frozen)
            self.assertEqual(reloaded.supersedes_id, unreviewed.id)
            self.assertIn("fecha_notificacion", reloaded.facts.unresolved)
            self.assertEqual(reloaded.facts.facts["fecha_notificacion"].sources, [])
            previous = get_validated_facts(conn, case_id, unreviewed.id)
            self.assertIsNotNone(previous.invalidated_at)


    def test_working_draft_versions_custody_staleness_and_atomicity(self):
        from rtm_core import traffic_working_draft as drafts
        import ops_operator_router as router
        from types import SimpleNamespace
        from fastapi import Request
        case_id, original_id, operator_id = (str(uuid.uuid4()) for _ in range(3))
        identity = {"full_name": "PRUEBA LOCAL", "dni_nie": "RTMTEST003",
                    "domicilio_notif": "Calle Ficticia 1", "local_test": drafts.MARKER}
        with self.engine.begin() as conn:
            conn.execute(text("""INSERT INTO cases(id,status,payment_status,authorized,department,case_type,interested_data,test_mode)
                VALUES (:id,'manual_review','paid',TRUE,'traffic','fine',CAST(:identity AS JSONB),TRUE)"""),
                {"id": case_id, "identity": json.dumps(identity)})
            conn.execute(text("INSERT INTO documents(id,case_id,kind,sha256,mime,size_bytes) VALUES (:id,:case,'original',:sha,'application/pdf',100)"),
                {"id": original_id, "case": case_id, "sha": "b"*64})
            facts = ValidatedFacts(case_id=case_id,service="traffic",extractor_version="synthetic_draft_test",
                source_document_ids=[original_id],unresolved=list(drafts.FIELDS),
                facts={key: {"status":"unresolved"} for key in drafts.FIELDS})
            initial = create_validated_facts(conn,case_id=case_id,facts=facts,created_by="operator:"+operator_id)
        scope=SimpleNamespace(individual_session=True,scope_all=True,role_code="rtm.supervisor",
                              permissions=("ops.view","ops.supervise"),operator_id=operator_id)
        request=Request({"type":"http"})
        request.state.rtm_operator_context=SimpleNamespace(operator_id=operator_id,
            session_id=str(uuid.uuid4()),actor="operator:"+operator_id)
        with tempfile.TemporaryDirectory(prefix="rtm-working-draft-test-") as directory, \
                patch.object(drafts,"local_operator_auth_requested",return_value=True), \
                patch.object(drafts.storage,"assert_local_document_storage_ready",return_value=__import__('pathlib').Path(directory)), \
                patch.object(router,"get_engine",return_value=self.engine), \
                patch.object(router,"require_operator_token"), \
                patch.object(router,"load_ops_case_scope",return_value=scope):
            state=router.get_working_draft(case_id,request,"test")["working_draft"]
            self.assertFalse(state["can_prepare"])
            self.assertEqual(len(state["blockers"]),8)
            self.assertIsNone(state["preparation_guide"])
            with self.engine.begin() as conn:
                seed_signed_case_authority(conn,case_id)
                values={"matricula":"RTM-TEST-003","organismo":"Ayuntamiento ficticio RTM",
                    "expediente_ref":"RTM-BORRADOR-003","fecha_documento":"2026-09-17",
                    "fecha_notificacion":"2026-09-18","sancion_importe_eur":200,
                    "hecho_denunciado_literal":"Estacionar en una zona de prueba señalizada."}
                corrected=review_facts(conn,case_id=case_id,facts_id=initial.id,actor="operator:"+operator_id,
                    body=ReviewFactsBody(expected_payload_sha256=initial.payload_sha256,reason="Datos sintéticos contrastados en integración",
                        changes=[dict(field=key,value=value,document_id=original_id,page_index=0,evidence=str(value)) for key,value in values.items()]))
                corrected=review_facts(conn,case_id=case_id,facts_id=corrected.id,actor="operator:"+operator_id,
                    body=ReviewFactsBody(expected_payload_sha256=corrected.payload_sha256,reason="Datos adicionales de prueba contrastados",
                        changes=[dict(operation="add",field=key,value=value,document_id=original_id,page_index=0,evidence=evidence)
                            for key,value,evidence in [("lugar_infraccion","Calle ficticia 1","Lugar: Calle ficticia 1"),
                                                      ("pago_multa_reducido",False,"Multa no pagada con reducción")]]))
                unchanged_case=conn.execute(text("SELECT to_jsonb(c) FROM cases c WHERE id=:id"),{"id":case_id}).scalar_one()
            state=router.get_working_draft(case_id,request,"test")["working_draft"]
            self.assertTrue(state["can_prepare"])
            self.assertEqual(state["preparation_guide"]["status"], "review_required")
            self.assertEqual(state["preparation_guide"]["case_id"], case_id)
            self.assertIn("no se ha pagado", state["preparation_guide"]["payment_note"])
            optional={item["key"]:item["value"] for item in state["reviewed_facts"]}
            self.assertEqual(optional["lugar_infraccion"],"Calle ficticia 1")
            self.assertIs(optional["pago_multa_reducido"],False)
            self.assertNotIn("fecha_limite",optional)
            def body(current, **changes):
                return drafts.WorkingDraftBody(expected_source_sha256=current["source_sha256"],
                    expected_latest_id=current["latest_id"],draft_acknowledged=True,
                    change_reason="Versión ficticia para comprobar persistencia",**changes)
            first=router.save_working_draft(case_id,body(state),request,"test")["working_draft"]
            self.assertEqual(first["source_sha256"],state["source_sha256"])
            self.assertTrue(first["history"][0]["current_source"])
            self.assertFalse(first["final_resource_generated"])
            self.assertEqual(first["history"][0]["preparation_guide"], state["preparation_guide"])
            first_id=first["latest_id"]
            response=router.get_working_draft_pdf(case_id,first_id,request,"test")
            rendered=" ".join(p.extract_text() for p in PdfReader(io.BytesIO(response.body)).pages)
            self.assertIn("RTM-BORRADOR-003",rendered)
            self.assertIn("Calle ficticia 1",rendered)
            self.assertIn("Multa pagada con reducción: No",rendered)
            self.assertIn("PENDIENTE DE REVISION JURIDICA",rendered)
            self.assertEqual(hashlib.sha256(response.body).hexdigest(),first["history"][0]["pdf"]["sha256"])
            self.assertNotIn("b2_key",json.dumps(first))
            with self.assertRaises(HTTPException) as conflict:
                router.save_working_draft(case_id,body(state),request,"test")
            self.assertEqual(conflict.exception.status_code,409)
            files_before=set(__import__('pathlib').Path(directory).rglob('*.pdf'))
            with patch.object(drafts,"_signed_envelope",side_effect=RuntimeError("audit failure")), self.assertRaisesRegex(RuntimeError,"audit failure"):
                router.save_working_draft(case_id,body(first),request,"test")
            self.assertEqual(set(__import__('pathlib').Path(directory).rglob('*.pdf')),files_before)
            self.assertEqual(router.get_working_draft(case_id,request,"test")["working_draft"]["latest_id"],first_id)
            second=router.save_working_draft(case_id,body(first,grounds="Texto sintético pendiente de contraste jurídico.",
                pending_notes=first["preparation_guide"]["pending_text"]),request,"test")["working_draft"]
            self.assertEqual(len(second["history"]),2)
            self.assertEqual(second["history"][0]["previous_id"],first_id)
            self.assertEqual(router.get_working_draft_pdf(case_id,first_id,request,"test").body,response.body)
            prepared_pdf = router.get_working_draft_pdf(case_id,second["latest_id"],request,"test").body
            self.assertIn("Precepto y ordenanza", " ".join(p.extract_text() for p in PdfReader(io.BytesIO(prepared_pdf)).pages))
            with patch.object(drafts, "build_parking_preparation", side_effect=lambda facts: {
                    **state["preparation_guide"], "message": "Regla de preparación revisada para la prueba"}):
                changed_guide=router.get_working_draft(case_id,request,"test")["working_draft"]
                self.assertNotEqual(changed_guide["source_sha256"], second["source_sha256"])
                self.assertFalse(changed_guide["history"][0]["current_source"])
                with self.assertRaises(HTTPException):router.save_working_draft(case_id,body(second),request,"test")
            with self.engine.begin() as conn:
                self.assertEqual(conn.execute(text("SELECT to_jsonb(c) FROM cases c WHERE id=:id"),{"id":case_id}).scalar_one(),unchanged_case)
                self.assertFalse(get_validated_facts(conn,case_id,corrected.id).frozen)
                self.assertEqual(conn.execute(text("SELECT count(*) FROM rtm_generated_resources WHERE case_id=:id"),{"id":case_id}).scalar_one(),0)
                conn.execute(text("UPDATE documents SET sha256=:sha WHERE id=:id"),{"id":original_id,"sha":"c"*64})
            stale=router.get_working_draft(case_id,request,"test")["working_draft"]
            self.assertFalse(stale["history"][0]["current_source"])
            with self.assertRaises(HTTPException): router.save_working_draft(case_id,body(second),request,"test")
            with self.engine.begin() as conn:
                conn.execute(text("UPDATE cases SET payment_status='unpaid' WHERE id=:id"),{"id":case_id})
            unpaid=router.get_working_draft(case_id,request,"test")["working_draft"]
            self.assertFalse(unpaid["can_prepare"])
            self.assertIsNone(unpaid["preparation_guide"])
            with self.engine.begin() as conn:
                conn.execute(text("UPDATE events SET payload=jsonb_set(payload,'{material,content}','\"altered\"'::jsonb) WHERE id=:id"),{"id":first_id})
            with self.assertRaises(HTTPException): router.get_working_draft_pdf(case_id,first_id,request,"test")
            with self.engine.begin() as conn:
                conn.execute(text("UPDATE cases SET test_mode=FALSE WHERE id=:id"),{"id":case_id})
            with self.assertRaises(HTTPException): router.get_working_draft(case_id,request,"test")


if __name__ == "__main__":
    unittest.main()
